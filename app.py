import os
import requests
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
from dotenv import load_dotenv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill

load_dotenv()
app = Flask(__name__)

SUPABASE_URL = os.environ.get('SUPABASE_URL')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY')

def sb(table):
    return f"{SUPABASE_URL}/rest/v1/{table}"

def hdrs():
    return {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "return=representation"
    }

def get(table, params=None):
    r = requests.get(sb(table), headers=hdrs(), params=params)
    return r.json()

def post(table, data):
    r = requests.post(sb(table), headers=hdrs(), json=data)
    return r.json()

def patch(table, match_params, data):
    r = requests.patch(sb(table), headers=hdrs(), params=match_params, json=data)
    return r.json()

def delete(table, match_params):
    r = requests.delete(sb(table), headers=hdrs(), params=match_params)
    return r.status_code


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/cities')
def get_cities():
    try:
        return jsonify(get('cities', {'order': 'name'}))
    except Exception as e:
        print(f"[get_cities] Error: {e}")
        return jsonify([])


@app.route('/api/classes')
def get_classes():
    try:
        params = {'order': 'name'}
        city_id = request.args.get('city_id')
        if city_id:
            params['city_id'] = f'eq.{city_id}'
        return jsonify(get('classes', params))
    except Exception as e:
        return jsonify([])


@app.route('/api/classes', methods=['POST'])
def add_class():
    try:
        data = request.json
        post('classes', {'city_id': data['city_id'], 'name': data['name']})
        _log('ADD_CLASS', f"Added class: {data['name']}", 'Admin')
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/tasks/<int:class_id>')
def get_tasks(class_id):
    try:
        return jsonify(get('tasks', {'class_id': f'eq.{class_id}', 'order': 'task_name'}))
    except Exception as e:
        return jsonify([])


@app.route('/api/tasks', methods=['POST'])
def add_task():
    try:
        data = request.json
        post('tasks', {
            'class_id':   data['class_id'],
            'task_name':  data['task_name'],
            'max_points': data.get('max_points', 100)
        })
        _log('ADD_TASK', f"Added task: {data['task_name']}", 'Admin')
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/students')
def get_students():
    try:
        city_id  = request.args.get('city_id')
        class_id = request.args.get('class_id')

        params = {'order': 'total_points.desc', 'select': '*,cities(name),classes(name)'}
        if city_id:
            params['city_id'] = f'eq.{city_id}'
        if class_id:
            params['class_id'] = f'eq.{class_id}'

        students = get('students', params)
        if not students:
            return jsonify([])

        student_ids = [s['id'] for s in students]
        class_ids   = list({s['class_id'] for s in students if s.get('class_id')})

        approvals_raw = get('approvals', {
            'student_id': f'in.({",".join(str(i) for i in student_ids)})',
            'select': 'student_id,task_id,grade_points'
        }) if student_ids else []

        tasks_raw = get('tasks', {
            'class_id': f'in.({",".join(str(i) for i in class_ids)})',
            'select': 'id,class_id,max_points'
        }) if class_ids else []

        task_max = {t['id']: t['max_points'] for t in tasks_raw}

        tasks_per_class = {}
        for t in tasks_raw:
            cid = t['class_id']
            tasks_per_class[cid] = tasks_per_class.get(cid, 0) + 1

        from collections import defaultdict
        student_task_points = defaultdict(lambda: defaultdict(int))
        for a in approvals_raw:
            student_task_points[a['student_id']][a['task_id']] += a['grade_points']

        completed_per_student = {}
        for sid, task_points in student_task_points.items():
            completed_per_student[sid] = sum(
                1 for tid, pts in task_points.items()
                if pts >= task_max.get(tid, 100)
            )

        result = []
        for s in students:
            result.append({
                'id':              s['id'],
                'name':            s['name'],
                'city_name':       s.get('cities', {}).get('name', 'Unknown') if s.get('cities') else 'Unknown',
                'class_name':      s.get('classes', {}).get('name', 'Not Assigned') if s.get('classes') else 'Not Assigned',
                'total_points':    s['total_points'],
                'tasks_completed': completed_per_student.get(s['id'], 0),
                'total_tasks':     tasks_per_class.get(s.get('class_id'), 0),
                'created_at':      s['created_at']
            })

        return jsonify(result)
    except Exception as e:
        import traceback; traceback.print_exc()
        return jsonify([])


@app.route('/api/students', methods=['POST'])
def add_student():
    try:
        data = request.json
        post('students', {
            'name':         data['name'],
            'city_id':      data['city_id'],
            'class_id':     data.get('class_id'),
            'total_points': 0
        })
        _log('ADD_STUDENT', f"Added student: {data['name']}", 'Admin')
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/students/<int:student_id>', methods=['GET'])
def get_student_details(student_id):
    try:
        students = get('students', {
            'id': f'eq.{student_id}',
            'select': '*,cities(name),classes(name)'
        })
        if not students:
            return jsonify({'student': None, 'approvals': []})

        s = students[0]
        student_data = {
            **s,
            'city_name':  s.get('cities', {}).get('name', 'Unknown') if s.get('cities') else 'Unknown',
            'class_name': s.get('classes', {}).get('name', 'Not Assigned') if s.get('classes') else 'Not Assigned',
        }

        approvals = get('approvals', {
            'student_id': f'eq.{student_id}',
            'select':     '*,tasks(task_name,max_points)',
            'order':      'approval_date.desc'
        })

        return jsonify({'student': student_data, 'approvals': approvals})
    except Exception as e:
        return jsonify({'student': None, 'approvals': []})


@app.route('/api/students/<int:student_id>', methods=['PUT'])
def update_student(student_id):
    try:
        data = request.json
        new_name = data.get('name', '').strip()
        if not new_name:
            return jsonify({'success': False, 'error': 'Name cannot be empty'}), 400
        patch('students', {'id': f'eq.{student_id}'}, {'name': new_name})
        _log('UPDATE_STUDENT', f"Renamed student #{student_id} to: {new_name}", 'Admin')
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/approvals', methods=['POST'])
def add_approval():
    try:
        data = request.json

        existing = get('students', {
            'name':    f'eq.{data["student_name"]}',
            'city_id': f'eq.{data["city_id"]}'
        })

        if not existing:
            new_student   = post('students', {
                'name':         data['student_name'],
                'city_id':      data['city_id'],
                'class_id':     data['class_id'],
                'total_points': 0
            })
            student_id    = new_student[0]['id']
            current_total = 0
        else:
            student_id    = existing[0]['id']
            current_total = existing[0]['total_points']

        post('approvals', {
            'student_id':       student_id,
            'class_id':         data['class_id'],
            'task_id':          data['task_id'],
            'grade_points':     data['grade_points'],
            'mentor_signature': data['mentor_signature'],
            'approval_date':    data['approval_date']
        })

        new_total = current_total + data['grade_points']
        patch('students', {'id': f'eq.{student_id}'}, {'total_points': new_total})

        _log('ADD_APPROVAL',
             f"Added {data['grade_points']} pts for {data['student_name']} (Total: {new_total})",
             data['mentor_signature'])

        return jsonify({'success': True, 'student_total': new_total})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/approvals/<int:approval_id>', methods=['DELETE'])
def delete_approval(approval_id):
    try:
        approval = get('approvals', {'id': f'eq.{approval_id}', 'select': 'student_id,grade_points'})
        if approval:
            student_id = approval[0]['student_id']
            points     = approval[0]['grade_points']
            delete('approvals', {'id': f'eq.{approval_id}'})
            student = get('students', {'id': f'eq.{student_id}', 'select': 'total_points'})
            if student:
                new_total = max(0, student[0]['total_points'] - points)
                patch('students', {'id': f'eq.{student_id}'}, {'total_points': new_total})
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/activity-log')
def get_activity_log():
    try:
        return jsonify(get('activity_log', {'order': 'timestamp.desc', 'limit': '50'}))
    except Exception as e:
        return jsonify([])


@app.route('/api/download-excel')
def download_excel():
    try:
        city_id = request.args.get('city_id')
        params  = {'select': '*,cities(name),classes(name)', 'order': 'total_points.desc'}
        if city_id:
            params['city_id'] = f'eq.{city_id}'
        students = get('students', params)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"
        for col, h in enumerate(['Student Name', 'City', 'Class', 'Total Points'], 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")

        for row, s in enumerate(students, 2):
            ws.cell(row=row, column=1, value=s['name'])
            ws.cell(row=row, column=2, value=s.get('cities', {}).get('name', '') if s.get('cities') else '')
            ws.cell(row=row, column=3, value=s.get('classes', {}).get('name', '') if s.get('classes') else '')
            ws.cell(row=row, column=4, value=s['total_points'])

        f = io.BytesIO()
        wb.save(f)
        f.seek(0)
        return send_file(f, as_attachment=True,
                         download_name=f'students_{datetime.now().strftime("%Y%m%d")}.xlsx',
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:
        return jsonify({'error': str(e)}), 500


def _log(action_type, description, user_name):
    try:
        post('activity_log', {
            'action_type': action_type,
            'description': description,
            'user_name':   user_name
        })
    except Exception as e:
        print(f"[_log] Warning: {e}")


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)